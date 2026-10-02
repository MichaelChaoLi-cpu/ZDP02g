from pathlib import Path
from copy import deepcopy
import json,hashlib,csv
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[2];out=root/'Rev/docs/r2c14-proposal';data=root/'data/results/revision_sample_descriptives'
source=root/'Rev/revision/ZDP02g.rev.clean.docx';d=Document(source);ref=d.tables[3];props=deepcopy(ref._tbl.tblPr)
for el in list(d._element.body):
 if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
for s in d.sections:
 for node in list(s._sectPr):
  if node.tag in [qn('w:lnNumType'),qn('w:headerReference'),qn('w:footerReference')]:s._sectPr.remove(node)
labels={
'GENDER':('Gender',{1:'Male',2:'Female',3:'Other'}),
'MARITAL_STATUS':('Marital status',{1:'Single / never married',2:'Married',3:'Separated',4:'Divorced',5:'Widowed',6:'Domestic partner'}),
'EMPLOYMENT':('Employment status',{1:'Employed for an employer',2:'Self-employed',3:'Retired',4:'Student',5:'Homemaker',6:'Unemployed and looking for a job',7:'None of these / other',8:'Code 8 (undefined in available codebook)'}),
'OWN_RENT_HOME_Y1':('Housing status (baseline)',{1:'Someone in household owns home',2:'Someone in household rents home',3:'Both',4:'Neither',5:'Rent',6:'Own',7:'Something else'}),
'URBAN_RURAL':('Residential area',{1:'Rural area or farm',2:'Small town or village',3:'Large city',4:'Suburb of a large city'}),
'EDUCATION_3':('Education',{1:'Elementary or less (up to 8 years)',2:'Secondary / some tertiary (9–15 years)',3:'Four years beyond high school / 4-year degree'}),
'HAVE_CHILD':('Child under 18 in household',{0:'No',1:'Yes'}),
'CLOSE_TO':('Knows a special person felt close to',{0:'No',1:'Yes'}),
'BODILY_PAIN':('Bodily pain in past four weeks',{0:'Not very much / none at all',1:'Some / a lot'}),
'HEALTH_PROB':('Activity-restricting health problems',{0:'No',1:'Yes'}),
'DAYS_EXERCISE':('Exercise ≥30 minutes in past week',{0:'No qualifying days',1:'At least one qualifying day'})}
freq=list(csv.DictReader((data/'categorical_frequencies.csv').open()));continuous=list(csv.DictReader((data/'continuous_summary.csv').open()));cw=list(csv.DictReader((data/'country_wave_counts.csv').open()))
notes=[]
def title(text):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1;r=p.add_run(text);r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(11)
def note(text):
 notes.append(text);p=d.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1;r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(9)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t._tbl.remove(t._tbl.tblPr);t._tbl.insert(0,deepcopy(props));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 allrows=[headers]+rows
 for i,values in enumerate(allrows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,(c,value) in enumerate(zip(cells,values)):
   c.width=Inches(widths[j]);mar=OxmlElement('w:tcMar')
   for edge in ['left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'35');e.set(qn('w:type'),'dxa');mar.append(e)
   c._tc.get_or_add_tcPr().append(mar);p=c.paragraphs[0];p.paragraph_format.left_indent=Pt(0);p.paragraph_format.right_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.RIGHT
   r=p.add_run(str(value));r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=i==0
  trpr=t.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
 return t
names=['Physical health (0–10)','Household income (annual nominal USD)','Latitude (°)','Longitude (°)','Age (years)','Adults in household (baseline)','Worry about expenses (0–10)','Annual mean temperature (°C)','Annual temperature SD (°C)']
title('Table 1. Data summary');title('Panel A. Continuous and count variables')
rows=[]
for name,r in zip(names,continuous):rows.append([name]+[(f'{float(r[k]):.3f}' if k in ['mean','std'] else f'{float(r[k]):.2f}') for k in ['mean','std','min','25%','50%','75%','max']])
table(['Variable','Mean','SD','Min','25%','Median','75%','Max'],rows,[1.23,.8,.8,.55,.62,.65,.6,.75])
note('All variables have 273,031 non-missing analytical records. SD is the sample standard deviation. Income is annual nominal USD; no purchasing-power-parity adjustment is applied. Temperature summaries use the corrected native-grid assignment.')
# Explicit page breaks keep each complete variable group together.
for ix,cols in enumerate([list(labels)[:4],list(labels)[4:]]):
 d.add_page_break();title('Table 1. Data summary (continued)');title('Panel B. Categorical variables'+(' (continued)' if ix else ''))
 rows=[]
 for col in cols:
  for r in [x for x in freq if x['variable']==col]:
   code=int(r['code']);pct=float(r['percent']);rows.append([labels[col][0]+' — '+labels[col][1][code],str(code),f"{int(r['n']):,}",'<0.01' if 0<pct<.01 else f'{pct:.2f}'])
 table(['Variable and category','Code','n','%'],rows,[4.15,.5,.75,.6])
 note('Percentages use all 273,031 analytical records and are unweighted. Totals may differ from 100% because of rounding. Category codes are retained to distinguish similarly named housing categories. Employment code 8 (n = 5) is not defined in the available baseline codebook and remains uninterpreted.')
 if ix:note('Housing status and household adult count use baseline values in both waves. Gender is a shared field across wave records. Labels follow the available baseline codebook and documented preprocessing; this does not independently verify later-wave questionnaire equivalence.')
d.add_page_break();title('Table 1. Data summary (continued)');title('Panel C. Analytical records by country/region and wave')
rows=[[r['Country'].removeprefix('the '),f"{int(r['Wave 1']):,}",f"{int(r['Wave 2']):,}",f"{int(r['Total']):,}",f"{100*int(r['Total'])/273031:.2f}"] for r in cw]
rows.append(['Total','175,507','97,524','273,031','100.00']);table(['Country / region','Wave 1','Wave 2','Total','% of pooled'],rows,[2.2,.95,.95,.95,.95])
note('Wave 1 contributes 175,507 records (64.28%) and Wave 2 contributes 97,524 records (35.72%). N refers to observation records, not unique respondents; people observed in both waves contribute two records. Wave labels do not imply that every interview occurred in the corresponding assigned exposure year.')
d.save(out/'Table1.candidate.docx')
(out/'category-labels.json').write_text(json.dumps(labels,indent=2,ensure_ascii=False))
(out/'artifact.md').write_text('Reference: '+str(source)+'\nSHA256: '+hashlib.sha256(source.read_bytes()).hexdigest()+'\nTable-only candidate. Retain Letter portrait, 1.25-inch side margins, Times New Roman, monochrome minimal table appearance. Clone source document and table properties; split panels at explicit page boundaries. Source manuscript remains immutable.\n')
(out/'proposal.md').write_text('# R2C14 complete table replacement proposal\n\nAction: Replace the complete current Table 1 (Data Summary) at its existing Tables-section location, before Table 2, with Table1.candidate.docx.\n\nScope: one table object, organized as Panels A–C over four review pages. Exact replacement text, labels, values and notes are contained in the candidate DOCX; all numerical cells derive from the validated aggregate CSVs. Continuous statistics retained; categorical mean/SD/quantiles replaced with n (%); country and wave summaries replaced with the cross-tabulation.\n\nBroader rewrite rationale: categorical codes do not have meaningful means/SDs; expanded categorical rows and country-by-wave counts require restructuring the table. No model or sample changes.\n\nPrior overlap: approved R2C3 corrected temperature rows; R2C15 annual nominal USD income label. Both are retained in the replacement. Protected objects: no citations, equations or other tables should change.\n\nProposed notes:\n\n'+'\n\n'.join(notes)+'\n\nStatus: candidate for exact-content review before tracked manuscript insertion.\n')
print('Saved',out/'Table1.candidate.docx')
